from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, jsonify
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from database import get_db_connection
import requests

app = Flask(__name__)

app.secret_key = "digital_farm_secret_key_123"


# ==========================================================
# HOME
# ==========================================================

@app.route("/")
def home():

    return render_template("index.html")


# ==========================================================
# REGISTER
# ==========================================================

@app.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        fullname = request.form.get("fullname", "").strip()
        mobile = request.form.get("mobile", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")

        connection = None
        cursor = None

        try:

            connection = get_db_connection()
            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO users
                (fullname, mobile, email, password)
                VALUES (%s, %s, %s, %s)
            """, (
                fullname,
                mobile,
                email,
                password
            ))

            connection.commit()

            flash(
                "Registration successful! Please login.",
                "success"
            )

            return redirect(url_for("login"))

        except Exception as e:

            if connection:
                connection.rollback()

            if "Duplicate entry" in str(e):

                flash(
                    "Email already registered. Please use another email.",
                    "error"
                )

                return render_template("register.html")

            return "Registration Error: " + str(e)

        finally:

            if cursor:
                cursor.close()

            if connection:
                connection.close()

    return render_template("register.html")


# ==========================================================
# LOGIN
# ==========================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")

        connection = None
        cursor = None

        try:

            connection = get_db_connection()

            cursor = connection.cursor(dictionary=True)

            cursor.execute("""
                SELECT *
                FROM users
                WHERE email = %s
                AND password = %s
            """, (
                email,
                password
            ))

            user = cursor.fetchone()

            if user:

                session["user_id"] = user["id"]
                session["fullname"] = user["fullname"]
                session["email"] = user["email"]

                return redirect(url_for("dashboard"))

            flash(
                "Invalid email or password.",
                "error"
            )

        except Exception as e:

            return "Login Error: " + str(e)

        finally:

            if cursor:
                cursor.close()

            if connection:
                connection.close()

    return render_template("login.html")


# ==========================================================
# LOGOUT
# ==========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("home"))


# ==========================================================
# DASHBOARD
# ==========================================================

@app.route("/dashboard")
def dashboard():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        cursor = connection.cursor()

        # Total crops

        cursor.execute("""
            SELECT COUNT(*)
            FROM crops
        """)

        total_crops = cursor.fetchone()[0]


        # Total land

        cursor.execute("""
            SELECT COALESCE(
                SUM(CAST(area AS DECIMAL(10,2))),
                0
            )
            FROM land
        """)

        total_land = cursor.fetchone()[0]


        # Total fertilizer cost

        cursor.execute("""
            SELECT COALESCE(
                SUM(cost),
                0
            )
            FROM fertilizer
        """)

        fertilizer_cost = cursor.fetchone()[0]


        # Irrigation records

        cursor.execute("""
            SELECT COUNT(*)
            FROM irrigation
        """)

        irrigation_records = cursor.fetchone()[0]


        # Total expenses

        cursor.execute("""
            SELECT COALESCE(
                SUM(amount),
                0
            )
            FROM expenses
        """)

        total_expenses = cursor.fetchone()[0]


        # Total income

        cursor.execute("""
            SELECT COALESCE(
                SUM(amount),
                0
            )
            FROM income
        """)

        total_income = cursor.fetchone()[0]


        # Total harvest

        cursor.execute("""
            SELECT COALESCE(
                SUM(
                    CAST(
                        yield_quantity AS DECIMAL(10,2)
                    )
                ),
                0
            )
            FROM harvest
        """)

        total_harvest = cursor.fetchone()[0]


        # Profit

        profit = float(total_income) - float(total_expenses)


        return render_template(
            "dashboard.html",

            farmer_name=session.get(
                "fullname",
                "Farmer"
            ),

            total_crops=total_crops,

            total_land=total_land,

            fertilizer_cost=fertilizer_cost,

            irrigation_records=irrigation_records,

            total_expenses=total_expenses,

            total_income=total_income,

            total_harvest=total_harvest,

            profit=profit
        )

    except Exception as e:

        return "Dashboard Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# CROP
# ==========================================================

@app.route("/crop", methods=["GET", "POST"])
def crop():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            crop_name = request.form.get(
                "crop_name",
                ""
            ).strip()

            season = request.form.get(
                "season",
                ""
            ).strip()

            area = request.form.get(
                "area",
                ""
            ).strip()

            sowing_date = request.form.get(
                "sowing_date",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO crops
                (
                    crop_name,
                    season,
                    area,
                    sowing_date
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                crop_name,
                season,
                area,
                sowing_date
            ))

            connection.commit()

            flash(
                "Crop record saved successfully!",
                "success"
            )

            return redirect(url_for("crop"))


        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute("""
            SELECT *
            FROM crops
            ORDER BY id DESC
        """)

        crops = cursor.fetchall()

        return render_template(
            "crop.html",
            crops=crops
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Crop Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# LAND
# ==========================================================

@app.route("/land", methods=["GET", "POST"])
def land():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            survey_no = request.form.get(
                "survey_no",
                ""
            ).strip()

            area = request.form.get(
                "area",
                ""
            ).strip()

            soil_type = request.form.get(
                "soil_type",
                ""
            ).strip()

            location = request.form.get(
                "location",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO land
                (
                    survey_no,
                    area,
                    soil_type,
                    location
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                survey_no,
                area,
                soil_type,
                location
            ))

            connection.commit()

            flash(
                "Land details saved successfully!",
                "success"
            )

            return redirect(url_for("land"))


        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute("""
            SELECT *
            FROM land
            ORDER BY id DESC
        """)

        lands = cursor.fetchall()

        return render_template(
            "land.html",
            lands=lands
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Land Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# FERTILIZER
# ==========================================================

@app.route("/fertilizer", methods=["GET", "POST"])
def fertilizer():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            fertilizer_name = request.form.get(
                "fertilizer_name",
                ""
            ).strip()

            quantity = request.form.get(
                "quantity",
                ""
            ).strip()

            cost = request.form.get(
                "cost",
                ""
            ).strip()

            date = request.form.get(
                "date",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO fertilizer
                (
                    fertilizer_name,
                    quantity,
                    cost,
                    date
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                fertilizer_name,
                quantity,
                cost,
                date
            ))

            connection.commit()

            flash(
                "Fertilizer record saved successfully!",
                "success"
            )

            return redirect(
                url_for("fertilizer")
            )


        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute("""
            SELECT *
            FROM fertilizer
            ORDER BY id DESC
        """)

        fertilizers = cursor.fetchall()

        return render_template(
            "fertilizer.html",
            fertilizers=fertilizers
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Fertilizer Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# IRRIGATION
# ==========================================================

@app.route("/irrigation", methods=["GET", "POST"])
def irrigation():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            water_source = request.form.get(
                "water_source",
                ""
            ).strip()

            method = request.form.get(
                "method",
                ""
            ).strip()

            date = request.form.get(
                "date",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO irrigation
                (
                    water_source,
                    method,
                    date
                )
                VALUES
                (
                    %s,
                    %s,
                    %s
                )
            """, (
                water_source,
                method,
                date
            ))

            connection.commit()

            flash(
                "Irrigation record saved successfully!",
                "success"
            )

            return redirect(
                url_for("irrigation")
            )


        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute("""
            SELECT *
            FROM irrigation
            ORDER BY id DESC
        """)

        irrigations = cursor.fetchall()

        return render_template(
            "irrigation.html",
            irrigations=irrigations
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Irrigation Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# EXPENSES
# ==========================================================

@app.route("/expenses", methods=["GET", "POST"])
def expenses():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            expense_type = request.form.get(
                "expense_type",
                ""
            ).strip()

            amount = request.form.get(
                "amount",
                ""
            ).strip()

            date = request.form.get(
                "date",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO expenses
                (
                    expense_type,
                    amount,
                    date
                )
                VALUES
                (
                    %s,
                    %s,
                    %s
                )
            """, (
                expense_type,
                amount,
                date
            ))

            connection.commit()

            flash(
                "Expense saved successfully!",
                "success"
            )

            return redirect(
                url_for("expenses")
            )


        cursor = connection.cursor(
            dictionary=True
        )

        cursor.execute("""
            SELECT *
            FROM expenses
            ORDER BY id DESC
        """)

        expenses_list = cursor.fetchall()

        return render_template(
            "expenses.html",
            expenses=expenses_list
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Expense Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# INCOME
# ==========================================================

@app.route("/income", methods=["GET", "POST"])
def income():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            crop_name = request.form.get(
                "crop_name",
                ""
            ).strip()

            quantity = request.form.get(
                "quantity",
                ""
            ).strip()

            amount = request.form.get(
                "amount",
                ""
            ).strip()

            date = request.form.get(
                "date",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO income
                (
                    crop_name,
                    quantity,
                    amount,
                    date
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                crop_name,
                quantity,
                amount,
                date
            ))

            connection.commit()

            flash(
                "Income saved successfully!",
                "success"
            )

            return redirect(
                url_for("income")
            )


        cursor = connection.cursor(
            dictionary=True
        )

        # Get income records

        cursor.execute("""
            SELECT *
            FROM income
            ORDER BY id DESC
        """)

        income_list = cursor.fetchall()


        # Get crop names for selection

        cursor.execute("""
            SELECT crop_name
            FROM crops
            ORDER BY crop_name
        """)

        crops = cursor.fetchall()


        return render_template(
            "income.html",
            income=income_list,

            incomes=income_list,
            crops=crops
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Income Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# HARVEST
# ==========================================================

@app.route("/harvest", methods=["GET", "POST"])
def harvest():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        if request.method == "POST":

            crop_name = request.form.get(
                "crop_name",
                ""
            ).strip()

            harvest_date = request.form.get(
                "harvest_date",
                ""
            ).strip()

            yield_quantity = request.form.get(
                "yield_quantity",
                ""
            ).strip()

            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO harvest
                (
                    crop_name,
                    harvest_date,
                    yield_quantity
                )
                VALUES
                (
                    %s,
                    %s,
                    %s
                )
            """, (
                crop_name,
                harvest_date,
                yield_quantity
            ))

            connection.commit()

            flash(
                "Harvest saved successfully!",
                "success"
            )

            return redirect(
                url_for("harvest")
            )


        cursor = connection.cursor(
            dictionary=True
        )

        # Harvest records

        cursor.execute("""
            SELECT *
            FROM harvest
            ORDER BY id DESC
        """)

        harvest_list = cursor.fetchall()


        # Crop names

        cursor.execute("""
            SELECT crop_name
            FROM crops
            ORDER BY crop_name
        """)

        crops = cursor.fetchall()


        return render_template(
            "harvest.html",
            harvest=harvest_list,
            crops=crops
        )

    except Exception as e:

        if connection:
            connection.rollback()

        return "Harvest Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# FARM REPORT
# ==========================================================

@app.route("/reports")
def reports():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:

        connection = get_db_connection()

        cursor = connection.cursor(
            dictionary=True
        )


        # --------------------------------------------------
        # TOTAL CROPS
        # --------------------------------------------------

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM crops
        """)

        total_crops = cursor.fetchone()["total"]


        # --------------------------------------------------
        # TOTAL LAND
        # --------------------------------------------------

        cursor.execute("""
            SELECT COALESCE(
                SUM(
                    CAST(
                        area AS DECIMAL(10,2)
                    )
                ),
                0
            ) AS total
            FROM land
        """)

        total_land = cursor.fetchone()["total"]


        # --------------------------------------------------
        # FERTILIZER COST
        # --------------------------------------------------

        cursor.execute("""
            SELECT COALESCE(
                SUM(cost),
                0
            ) AS total
            FROM fertilizer
        """)

        fertilizer_cost = cursor.fetchone()["total"]


        # --------------------------------------------------
        # IRRIGATION RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM irrigation
        """)

        irrigation_records = cursor.fetchone()["total"]


        # --------------------------------------------------
        # TOTAL EXPENSES
        # --------------------------------------------------

        cursor.execute("""
            SELECT COALESCE(
                SUM(amount),
                0
            ) AS total
            FROM expenses
        """)

        total_expenses = cursor.fetchone()["total"]


        # --------------------------------------------------
        # TOTAL INCOME
        # --------------------------------------------------

        cursor.execute("""
            SELECT COALESCE(
                SUM(amount),
                0
            ) AS total
            FROM income
        """)

        total_income = cursor.fetchone()["total"]


        # --------------------------------------------------
        # TOTAL HARVEST
        # --------------------------------------------------

        cursor.execute("""
            SELECT COALESCE(
                SUM(
                    CAST(
                        yield_quantity AS DECIMAL(10,2)
                    )
                ),
                0
            ) AS total
            FROM harvest
        """)

        total_harvest = cursor.fetchone()["total"]


        # --------------------------------------------------
        # PROFIT
        # --------------------------------------------------

        profit = (
            float(total_income)
            -
            float(total_expenses)
        )


        # --------------------------------------------------
        # ALL CROP RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM crops
            ORDER BY id DESC
        """)

        crops = cursor.fetchall()


        # --------------------------------------------------
        # ALL LAND RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM land
            ORDER BY id DESC
        """)

        lands = cursor.fetchall()


        # --------------------------------------------------
        # ALL FERTILIZER RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM fertilizer
            ORDER BY id DESC
        """)

        fertilizers = cursor.fetchall()


        # --------------------------------------------------
        # ALL IRRIGATION RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM irrigation
            ORDER BY id DESC
        """)

        irrigations = cursor.fetchall()


        # --------------------------------------------------
        # ALL EXPENSE RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM expenses
            ORDER BY id DESC
        """)

        expenses_list = cursor.fetchall()


        # --------------------------------------------------
        # ALL INCOME RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM income
            ORDER BY id DESC
        """)

        income_list = cursor.fetchall()


        # --------------------------------------------------
        # ALL HARVEST RECORDS
        # --------------------------------------------------

        cursor.execute("""
            SELECT *
            FROM harvest
            ORDER BY id DESC
        """)

        harvest_list = cursor.fetchall()


        return render_template(

            "reports.html",

            farmer_name=session.get(
                "fullname",
                "Farmer"
            ),

            total_crops=total_crops,

            total_land=total_land,

            fertilizer_cost=fertilizer_cost,

            irrigation_records=irrigation_records,

            total_expenses=total_expenses,

            total_income=total_income,

            total_harvest=total_harvest,

            profit=profit,

            crops=crops,

            lands=lands,

            land_records=lands,

            fertilizers=fertilizers,

            irrigations=irrigations,

            expenses=expenses_list,

            income=income_list,

            incomes=income_list,

            harvest=harvest_list
        )

    except Exception as e:

        return "Reports Error: " + str(e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# DOWNLOAD PDF REPORT
# ==========================================================

@app.route("/download/pdf")
def download_pdf():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("SELECT COUNT(*) AS total FROM crops")
        total_crops = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COALESCE(SUM(CAST(area AS DECIMAL(10,2))), 0) AS total
            FROM land
        """)
        total_land = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(cost), 0) AS total FROM fertilizer")
        fertilizer_cost = cursor.fetchone()["total"]

        cursor.execute("SELECT COUNT(*) AS total FROM irrigation")
        irrigation_records = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM expenses")
        total_expenses = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM income")
        total_income = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COALESCE(
                SUM(CAST(yield_quantity AS DECIMAL(10,2))), 0
            ) AS total
            FROM harvest
        """)
        total_harvest = cursor.fetchone()["total"]

        profit = float(total_income) - float(total_expenses)

        cursor.execute("SELECT * FROM crops ORDER BY id DESC")
        crops = cursor.fetchall()

        cursor.execute("SELECT * FROM land ORDER BY id DESC")
        lands = cursor.fetchall()

        cursor.execute("SELECT * FROM fertilizer ORDER BY id DESC")
        fertilizers = cursor.fetchall()

        cursor.execute("SELECT * FROM irrigation ORDER BY id DESC")
        irrigations = cursor.fetchall()

        cursor.execute("SELECT * FROM expenses ORDER BY id DESC")
        expenses_list = cursor.fetchall()

        cursor.execute("SELECT * FROM income ORDER BY id DESC")
        income_list = cursor.fetchall()

        cursor.execute("SELECT * FROM harvest ORDER BY id DESC")
        harvest_list = cursor.fetchall()

        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "FarmTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=20,
            spaceAfter=10
        )

        story = []

        story.append(Paragraph("DIGITAL FARM RECORD SYSTEM", title_style))
        story.append(Paragraph(
            "Complete Farm Report - " +
            str(session.get("fullname", "Farmer")),
            styles["Heading2"]
        ))
        story.append(Spacer(1, 12))

        summary_data = [
            ["Total Crops", "Total Land", "Fertilizer Cost", "Irrigation Records",
             "Total Expenses", "Total Income", "Total Harvest", "Profit"],
            [str(total_crops), str(total_land), str(fertilizer_cost),
             str(irrigation_records), str(total_expenses), str(total_income),
             str(total_harvest), str(profit)]
        ]

        summary_table = Table(summary_data, repeatRows=1)
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e7d32")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 1), (-1, 1), colors.whitesmoke),
            ("PADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 18))

        def add_section(title, rows, columns):
            story.append(Paragraph(title, styles["Heading2"]))
            data = [columns]
            for row in rows:
                data.append([str(row.get(column, "")) for column in columns])

            if len(data) == 1:
                data.append(["No records available"] + [""] * (len(columns) - 1))

            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#388e3c")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(table)
            story.append(Spacer(1, 15))

        add_section("Crop Records", crops,
                    ["id", "crop_name", "season", "area", "sowing_date"])
        add_section("Land Records", lands,
                    ["id", "survey_no", "area", "soil_type", "location"])
        add_section("Fertilizer Records", fertilizers,
                    ["id", "fertilizer_name", "quantity", "cost", "date"])
        add_section("Irrigation Records", irrigations,
                    ["id", "water_source", "method", "date"])
        add_section("Expense Records", expenses_list,
                    ["id", "expense_type", "amount", "date"])
        add_section("Income Records", income_list,
                    ["id", "crop_name", "quantity", "amount", "date"])
        add_section("Harvest Records", harvest_list,
                    ["id", "crop_name", "harvest_date", "yield_quantity"])

        document.build(story)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="Digital_Farm_Report.pdf",
            mimetype="application/pdf"
        )

    except Exception as e:
        return "PDF Download Error: " + str(e)

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


# ==========================================================
# DOWNLOAD EXCEL REPORT
# ==========================================================

@app.route("/download/excel")
def download_excel():

    if "user_id" not in session:
        return redirect(url_for("login"))

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        queries = {
            "Crops": "SELECT * FROM crops ORDER BY id DESC",
            "Land": "SELECT * FROM land ORDER BY id DESC",
            "Fertilizer": "SELECT * FROM fertilizer ORDER BY id DESC",
            "Irrigation": "SELECT * FROM irrigation ORDER BY id DESC",
            "Expenses": "SELECT * FROM expenses ORDER BY id DESC",
            "Income": "SELECT * FROM income ORDER BY id DESC",
            "Harvest": "SELECT * FROM harvest ORDER BY id DESC"
        }

        records = {}
        for sheet_name, query in queries.items():
            cursor.execute(query)
            records[sheet_name] = cursor.fetchall()

        cursor.execute("SELECT COUNT(*) AS total FROM crops")
        total_crops = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COALESCE(SUM(CAST(area AS DECIMAL(10,2))), 0) AS total
            FROM land
        """)
        total_land = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(cost), 0) AS total FROM fertilizer")
        fertilizer_cost = cursor.fetchone()["total"]

        cursor.execute("SELECT COUNT(*) AS total FROM irrigation")
        irrigation_records = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM expenses")
        total_expenses = cursor.fetchone()["total"]

        cursor.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM income")
        total_income = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COALESCE(
                SUM(CAST(yield_quantity AS DECIMAL(10,2))), 0
            ) AS total
            FROM harvest
        """)
        total_harvest = cursor.fetchone()["total"]

        profit = float(total_income) - float(total_expenses)

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        summary = workbook.create_sheet("Farm Summary")
        summary.append(["DIGITAL FARM RECORD SYSTEM"])
        summary.append(["Farmer Name", session.get("fullname", "Farmer")])
        summary.append([])
        summary.append(["Metric", "Value"])
        summary.append(["Total Crops", total_crops])
        summary.append(["Total Land (Acres)", total_land])
        summary.append(["Fertilizer Cost", fertilizer_cost])
        summary.append(["Irrigation Records", irrigation_records])
        summary.append(["Total Expenses", total_expenses])
        summary.append(["Total Income", total_income])
        summary.append(["Total Harvest", total_harvest])
        summary.append(["Profit", profit])

        for cell in summary[1]:
            cell.font = Font(bold=True, size=16)

        for cell in summary[4]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="2E7D32")
            cell.font = Font(bold=True, color="FFFFFF")

        for row in summary.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")

        summary.column_dimensions["A"].width = 28
        summary.column_dimensions["B"].width = 25

        for sheet_name, rows in records.items():
            sheet = workbook.create_sheet(sheet_name)

            if rows:
                headers = list(rows[0].keys())
                sheet.append(headers)

                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="2E7D32")
                    cell.alignment = Alignment(horizontal="center")

                for row in rows:
                    sheet.append([row.get(header, "") for header in headers])

                for column_cells in sheet.columns:
                    max_length = 0
                    column_letter = column_cells[0].column_letter
                    for cell in column_cells:
                        value = "" if cell.value is None else str(cell.value)
                        max_length = max(max_length, len(value))
                    sheet.column_dimensions[column_letter].width = min(max_length + 3, 40)
            else:
                sheet.append(["No records available"])
                sheet.column_dimensions["A"].width = 25

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="Digital_Farm_Report.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        return "Excel Download Error: " + str(e)

    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


# ==========================================================
# WEATHER MODULE - OPENWEATHER
# ==========================================================

# Paste your OpenWeather API key between the quotes below.
# Example:
# OPENWEATHER_API_KEY = "your_actual_api_key"
OPENWEATHER_API_KEY = "YOUR_API_KEY"


@app.route("/weather")
def weather():

    if "user_id" not in session:
        return redirect(url_for("login"))

    return render_template("weather.html")


@app.route("/api/weather")
def api_weather():

    if "user_id" not in session:
        return jsonify({
            "success": False,
            "error": "Please login first."
        }), 401

    city = request.args.get("city", "").strip()

    if not city:
        return jsonify({
            "success": False,
            "error": "Please enter a place or city name."
        }), 400

    if OPENWEATHER_API_KEY == "YOUR_OPENWEATHER_API_KEY_HERE":
        return jsonify({
            "success": False,
            "error": "OpenWeather API key is not configured in app.py."
        }), 500

    try:

        # --------------------------------------------------
        # STEP 1: Convert place name to latitude/longitude
        # --------------------------------------------------
        geocode_response = requests.get(
            "https://api.openweathermap.org/geo/1.0/direct",
            params={
                "q": city,
                "limit": 1,
                "appid": OPENWEATHER_API_KEY
            },
            timeout=10
        )

        if geocode_response.status_code != 200:
            return jsonify({
                "success": False,
                "error": "Unable to find this location."
            }), geocode_response.status_code

        locations = geocode_response.json()

        if not locations:
            return jsonify({
                "success": False,
                "error": "Location not found. Please enter a valid city or place."
            }), 404

        location = locations[0]

        latitude = location.get("lat")
        longitude = location.get("lon")
        location_name = location.get("name", city)
        country = location.get("country", "")
        state = location.get("state", "")

        # --------------------------------------------------
        # STEP 2: Get current weather using coordinates
        # --------------------------------------------------
        weather_response = requests.get(
            "https://api.openweathermap.org/data/2.5/weather",
            params={
                "lat": latitude,
                "lon": longitude,
                "appid": OPENWEATHER_API_KEY,
                "units": "metric"
            },
            timeout=10
        )

        if weather_response.status_code != 200:

            try:
                error_data = weather_response.json()
                error_message = error_data.get(
                    "message",
                    "Weather service error."
                )
            except ValueError:
                error_message = "Weather service returned an invalid response."

            return jsonify({
                "success": False,
                "error": error_message
            }), weather_response.status_code

        data = weather_response.json()

        weather_info = data.get("weather", [{}])[0]
        main_info = data.get("main", {})
        wind_info = data.get("wind", {})
        clouds_info = data.get("clouds", {})
        sys_info = data.get("sys", {})

        return jsonify({
            "success": True,
            "location": location_name,
            "state": state,
            "country": country,
            "latitude": latitude,
            "longitude": longitude,
            "temperature": main_info.get("temp"),
            "feels_like": main_info.get("feels_like"),
            "min_temperature": main_info.get("temp_min"),
            "max_temperature": main_info.get("temp_max"),
            "humidity": main_info.get("humidity"),
            "pressure": main_info.get("pressure"),
            "visibility": data.get("visibility"),
            "wind_speed": wind_info.get("speed"),
            "wind_direction": wind_info.get("deg"),
            "cloudiness": clouds_info.get("all"),
            "description": weather_info.get("description", ""),
            "condition": weather_info.get("main", ""),
            "icon": weather_info.get("icon", ""),
            "sunrise": sys_info.get("sunrise"),
            "sunset": sys_info.get("sunset")
        })

    except requests.exceptions.Timeout:

        return jsonify({
            "success": False,
            "error": "Weather service took too long to respond. Please try again."
        }), 504

    except requests.exceptions.ConnectionError:

        return jsonify({
            "success": False,
            "error": "Could not connect to the weather service. Check your internet connection."
        }), 502

    except requests.exceptions.RequestException as e:

        return jsonify({
            "success": False,
            "error": "Weather service connection error: " + str(e)
        }), 502

    except ValueError:

        return jsonify({
            "success": False,
            "error": "Weather service returned invalid JSON data."
        }), 502

    except Exception as e:

        return jsonify({
            "success": False,
            "error": "Weather error: " + str(e)
        }), 500


# ==========================================================
# RUN APPLICATION
# ==========================================================

if __name__ == "__main__":

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )